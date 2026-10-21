"""
API tests for /api/extract and /api/export (Flask test client, no server).

    .venv/Scripts/python -m pytest tests -q
"""

from __future__ import annotations

import io
import json
import os
import sys

import cv2
import openpyxl
import pytest
from flask import Flask

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "backend"))
sys.path.insert(0, HERE)

import extract_routes  # noqa: E402
from synth import BoxSpec, ChartSpec, render, render_box  # noqa: E402

SPEC = ChartSpec([12.0, 31.5, 22.0], [1.5, 3.0, 2.2], labels=["Control", "CUMS", "CUMS+EA"])


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    extract_routes.JOBS_DIR = str(tmp_path_factory.mktemp("jobs"))
    app = Flask(__name__)
    app.add_url_rule("/api/extract", "extract", extract_routes.handle_extract, methods=["POST"])
    app.add_url_rule("/api/export", "export", extract_routes.handle_export, methods=["POST"])
    app.add_url_rule("/static/jobs/<job>/<filename>", "job_file", extract_routes.handle_job_file)
    return app.test_client()


def _png_bytes(spec=SPEC) -> bytes:
    ok, buf = cv2.imencode(".png", render(spec))
    assert ok
    return buf.tobytes()


def _pdf_bytes(caption: str) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(60, 60, 420, 400), stream=_png_bytes())
    page.insert_text((60, 430), caption, fontsize=9)
    return doc.tobytes()


def _post(client, data: bytes, name: str, **form):
    payload = {"file": (io.BytesIO(data), name)}
    payload.update(form)
    return client.post("/api/extract", data=payload, content_type="multipart/form-data")


def test_extract_png(client):
    r = _post(client, _png_bytes(), "chart.png")
    assert r.status_code == 200, r.get_json()
    data = r.get_json()
    assert len(data["figures"]) == 1
    panel = data["figures"][0]["panels"][0]
    assert [round(b["mean"], 1) for b in panel["bars"]] == pytest.approx([12.0, 31.5, 22.0], abs=0.4)
    assert [b["label"] for b in panel["bars"]] == ["Control", "CUMS", "CUMS+EA"]
    assert panel["error_type"] == "sem"
    # image is served back
    img = client.get(data["figures"][0]["image_url"])
    assert img.status_code == 200 and len(img.data) > 1000


def test_extract_pdf_reads_caption(client):
    r = _post(client, _pdf_bytes("Fig. 3. Behavioural tests. Data are mean ± SD (n = 8)."), "paper.pdf")
    assert r.status_code == 200, r.get_json()
    fig = r.get_json()["figures"][0]
    assert fig["figure"] == "3"
    assert fig["error_type_from_caption"] == "sd"
    assert fig["panels"][0]["error_type"] == "sd"


def test_error_type_override_and_options(client):
    r = _post(client, _png_bytes(), "chart.png", error_type="sd",
              options=json.dumps({"bar_names": False, "titles": False}))
    panel = r.get_json()["figures"][0]["panels"][0]
    assert panel["error_type"] == "sd"
    assert all(b["label"] == "" for b in panel["bars"])
    assert panel["y_title"] == ""


@pytest.mark.parametrize("form,name,data,status", [
    ({}, None, None, 400),                                     # no file
    ({}, "notes.txt", b"hello", 400),                          # unsupported type
    ({"options": "{not json"}, "chart.png", None, 400),        # bad options
    ({"error_type": "variance"}, "chart.png", None, 400),      # bad error type
    ({"job": "../../etc"}, "chart.png", None, 400),            # path tricks in job id
])
def test_extract_rejects_bad_input(client, form, name, data, status):
    if name is None:
        r = client.post("/api/extract", data={}, content_type="multipart/form-data")
    else:
        r = _post(client, data if data is not None else _png_bytes(), name, **form)
    assert r.status_code == status
    assert "error" in r.get_json()


def test_extract_corrupt_files(client):
    r = _post(client, b"%PDF-1.4 garbage", "broken.pdf")
    assert r.status_code in (400, 500) and "error" in r.get_json()
    r = _post(client, b"\x89PNG not really", "broken.png")
    assert r.status_code == 200 and r.get_json()["figures"] == []


def test_export_roundtrip(client):
    data = _post(client, _png_bytes(), "chart.png").get_json()
    fig = data["figures"][0]
    panel = fig["panels"][0]
    panel["error_type"] = "sd"
    for b in panel["bars"]:
        b["n"] = 8
    panel["bars"][0]["label"] = "control"
    r = client.post("/api/export", json={"figures": [fig], "filename": "x.xlsx"})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames[:2] == ["Summary", "All data"]
    ws = wb[wb.sheetnames[2]]
    header = [c.value for c in ws[1]][:12]
    assert header[0] == "mean" and header[-1] == "mean+sd" and header[5] == "treatment_group"
    row = [c.value for c in ws[2]][:12]
    assert row[5] == "control" and row[6] == 8
    assert abs(row[0] - 12.0) < 0.4
    assert isinstance(row[1], float) and abs(row[1] - 1.5) < 0.5       # sd = measured error
    assert str(row[10]).startswith("=")                              # sem derived by formula
    assert abs(row[11] - (row[0] + row[1])) < 1e-9


def test_box_plot_extract_and_export(client):
    """A box plot comes back as kind "box"; the export estimates mean/SD with live formulas."""
    stats = [(2, 5, 7, 9, 13), (4, 8, 11, 15, 19)]
    ok, buf = cv2.imencode(".png", render_box(BoxSpec(stats, ylim=(0, 20))))
    data = _post(client, buf.tobytes(), "box.png").get_json()
    panel = data["figures"][0]["panels"][0]
    assert panel["kind"] == "box" and len(panel["boxes"]) == 2
    assert abs(panel["boxes"][0]["median"] - 7) < 0.2
    panel["boxes"][0]["n"] = 12
    panel["whisker_type"] = "tukey"
    panel["mean_method"] = "not a method"          # falls back to the default
    r = client.post("/api/export", json={"figures": data["figures"]})
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb[wb.sheetnames[2]]
    header = [c.value for c in ws[1]]
    assert header[11] == "mean+sd" and header[14:19] == ["min", "Q1", "median", "Q3", "max"]
    row = [c.value for c in ws[2]]
    assert row[6] == 12
    assert str(row[0]).startswith('=IF(G2="","",') and "NORMSINV" in str(row[1])
    assert abs(row[16] - 7) < 0.2 and row[19] == "tukey" and row[20] == "luo"
    assert [c.value for c in wb["Summary"][2]][4] == "box"
    combined = [c.value for c in wb["All data"][1]]
    assert combined[-8:-3] == ["min", "Q1", "median", "Q3", "max"]


def test_export_excluded_and_empty(client):
    data = _post(client, _png_bytes(), "chart.png").get_json()
    fig = data["figures"][0]
    fig["panels"][0]["excluded"] = True
    assert client.post("/api/export", json={"figures": [fig]}).status_code == 400
    assert client.post("/api/export", json={}).status_code == 400


@pytest.mark.parametrize("bad", [
    {"figures": [{"panels": [{"bars": [{"x1": 0, "x2": 1, "top_px": 1, "cap_px": 1, "mean": "abc"}]}]}]},
    {"figures": [{"panels": ["not a panel"]}]},
    {"figures": [{"panels": [{"bars": [{"x1": 0, "x2": 1, "top_px": 1, "cap_px": 1, "mean": 1.0, "n": "eight"}]}]}]},
    {"figures": "nope"},
    {"figures": [{"panels": [{"kind": "box", "boxes": [{"x1": 0, "x2": 1, "q1_px": 1, "median_px": 1, "q3_px": 1,
                                                         "lo_px": 1, "hi_px": 1, "median": "x"}]}]}]},
])
def test_export_rejects_malformed(client, bad):
    r = client.post("/api/export", json=bad)
    assert r.status_code == 400 and "error" in r.get_json()


def test_old_jobs_cleaned(client, tmp_path):
    old = os.path.join(extract_routes.JOBS_DIR, "deadbeef")
    os.makedirs(old, exist_ok=True)
    open(os.path.join(old, "a.jpg"), "wb").close()
    os.utime(old, (0, 0))
    _post(client, _png_bytes(), "chart.png")
    assert not os.path.exists(old)
