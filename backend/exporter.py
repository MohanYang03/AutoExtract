"""
Spreadsheet export in the meta-analysis "long format".

Each panel gets its own sheet whose first 12 columns match the hand-made
Long_format_*.xlsx files exactly, so rows can be pasted straight across.
`sd` is a live formula (sem * sqrt(sample_size)) that fills in once the
sample size is typed in. Bars whose error bar could not be confirmed are
highlighted for a quick manual check.

Box plots fill the same columns with an estimated mean and SD: live
formulas on the five-number summary (written next to them) and the sample
size, using the method chosen for the panel (see estimators.py).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from types import SimpleNamespace
from typing import List, Optional

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from estimators import skewness_warning
from figure_extractor import Panel

LONG_COLUMNS = ["mean", "sd", "age(week)", "sex(male)", "acupoint", "treatment_group",
                "sample_size", "treatment_duration", "depression_model_duration",
                "paper_id", "sem", "mean+sem"]

FLAG_FILL = PatternFill("solid", start_color="FFF2CC")
HEADER_FONT = Font(bold=True)


@dataclass
class PanelResult:
    source: str            # file name the figure came from
    page: int              # 0-based page index (or -1 for an image file)
    figure: str            # figure number from the caption, if found
    panel: Panel
    paper_id: Optional[int] = None

    @property
    def name(self) -> str:
        if self.figure:
            fig = f"Fig{self.figure}"
        elif self.page >= 0:
            fig = f"p{self.page + 1}"
        else:
            fig = "Image"
        return f"{fig}{self.panel.letter}" if self.panel.letter else fig


def paper_id_from_path(path: str) -> Optional[int]:
    m = re.search(r"\(paper(\d+)\)", path)
    return int(m.group(1)) if m else None


def _sheet_title(base: str, used: set) -> str:
    title = re.sub(r"[\[\]\*\?/\\:]", "", base)[:31] or "Panel"
    candidate, i = title, 2
    while candidate in used:
        suffix = f" ({i})"
        candidate = title[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


BOX_COLUMNS = ["min", "Q1", "median", "Q3", "max", "whiskers", "mean method", "sd method"]


def long_header(error_type: str) -> List[str]:
    """LONG_COLUMNS with the last column named after what the error bars show."""
    return LONG_COLUMNS[:-1] + ["mean+sd" if error_type == "sd" else "mean+sem"]


def box_formulas(a: str, q1: str, m: str, q3: str, b: str, n: str,
                 whisker_type: str, mean_method: str, sd_method: str, quartiles: bool = True):
    """
    Excel formulas (without "=") for the estimated mean and SD, mirroring
    estimators.estimate. Arguments are cell references. When the whiskers
    are the range the five-number estimators are used (S2), otherwise the
    quartile ones (S3).
    """
    xi = f"(2*NORMSINV(({n}-0.375)/({n}+0.25)))"
    eta = f"(2*NORMSINV((0.75*{n}-0.125)/({n}+0.25)))"
    hozo_sd = (f"IF({n}<=15,SQRT((({a}-2*{m}+{b})^2/4+({b}-{a})^2)/12),"
               f"IF({n}<=70,({b}-{a})/4,({b}-{a})/6))")
    if not quartiles:
        # S1: min, median, max (a violin showing only its median)
        if mean_method == "luo":
            w = f"(4/(4+{n}^0.75))"
            mean = f"{w}*({a}+{b})/2+(1-{w})*{m}"
        else:
            mean = f"({a}+2*{m}+{b})/4"
        sd = hozo_sd if sd_method == "hozo" else f"({b}-{a})/{xi}"
        return mean, sd
    if whisker_type == "minmax":
        if mean_method == "wan":
            mean = f"({a}+2*{q1}+2*{m}+2*{q3}+{b})/8"
        elif mean_method == "hozo":
            mean = f"({a}+2*{m}+{b})/4"
        else:
            w1 = f"(2.2/(2.2+{n}^0.75))"
            w2 = f"(0.7-0.72/{n}^0.55)"
            mean = f"{w1}*({a}+{b})/2+{w2}*({q1}+{q3})/2+(1-{w1}-{w2})*{m}"
        if sd_method == "wan":
            sd = f"0.5*(({b}-{a})/{xi}+({q3}-{q1})/{eta})"
        elif sd_method == "hozo":
            sd = hozo_sd
        else:
            w = f"(1/(1+0.07*{n}^0.6))"
            sd = f"{w}*({b}-{a})/{xi}+(1-{w})*({q3}-{q1})/{eta}"
    else:
        if mean_method in ("wan", "hozo"):
            mean = f"({q1}+{m}+{q3})/3"
        else:
            w = f"(0.7+0.39/{n})"
            mean = f"{w}*({q1}+{q3})/2+(1-{w})*{m}"
        sd = f"({q3}-{q1})/{eta}"
    return mean, sd


def _write_box_rows(ws, start_row: int, result: PanelResult, first_col: int, stat_col: int) -> int:
    """
    One row per box: estimated mean/SD (formulas on n) in the long-format
    columns, and the five-number summary from `stat_col` on.
    """
    p = result.panel
    col = {name: first_col + i for i, name in enumerate(LONG_COLUMNS)}
    r = start_row
    for bx in p.boxes:
        refs = [f"{get_column_letter(stat_col + i)}{r}" for i in range(5)]
        n_ref = f"{get_column_letter(col['sample_size'])}{r}"
        mean_ref = f"{get_column_letter(col['mean'])}{r}"
        sd_ref = f"{get_column_letter(col['sd'])}{r}"
        five = [bx.lo, bx.q1, bx.median, bx.q3, bx.hi]
        for i, v in enumerate(five + [p.whisker_type, p.mean_method, p.sd_method]):
            ws.cell(r, stat_col + i, v)
        quartiles = bx.q1 is not None and bx.q3 is not None
        usable = (None not in five) if quartiles else (
            p.whisker_type == "minmax" and None not in (bx.lo, bx.median, bx.hi))
        if usable:
            mean_f, sd_f = box_formulas(*refs, n_ref, p.whisker_type, p.mean_method, p.sd_method,
                                        quartiles=quartiles)
            ws.cell(r, col["mean"], f'=IF({n_ref}="","",{mean_f})')
            ws.cell(r, col["sd"], f'=IF({n_ref}="","",{sd_f})')
            ws.cell(r, col["sem"], f'=IF({n_ref}="","",{sd_ref}/SQRT({n_ref}))')
            ws.cell(r, col["mean+sem"], f'=IF({n_ref}="","",{mean_ref}+{sd_ref})')
        ws.cell(r, col["treatment_group"], bx.label)
        if bx.n:
            ws.cell(r, col["sample_size"], bx.n)
        if result.paper_id is not None:
            ws.cell(r, col["paper_id"], result.paper_id)
        if bx.uncertain:
            for c in range(first_col, first_col + len(LONG_COLUMNS)):
                ws.cell(r, c).fill = FLAG_FILL
        r += 1
    return r


def _box_check(bx) -> str:
    notes = []
    if bx.uncertain:
        notes.append(bx.note or "check box")
    elif bx.note:
        notes.append(bx.note)
    if None not in (bx.q1, bx.median, bx.q3) and skewness_warning(bx.q1, bx.median, bx.q3):
        notes.append("skewed: estimate may be biased")
    return "; ".join(notes)


def line_rows(panel: Panel) -> list:
    """
    A line plot's points as long-format rows: the group is the series, and
    the x value (or x label) of each point goes alongside.
    """
    rows = []
    for k, s in enumerate(panel.series, 1):
        name = s.label or f"Series {k}"
        for pt in s.points:
            x = pt.x if pt.x is not None else pt.x_label
            if isinstance(x, float):
                x = round(x, 4)
            rows.append(SimpleNamespace(mean=pt.mean, mean_plus_err=pt.mean_plus_err, label=name, n=pt.n,
                                        uncertain=pt.uncertain, x=x))
    return rows


def _write_long_rows(ws, start_row: int, result: PanelResult, first_col: int = 1, items=None) -> int:
    """
    Write one row per bar in LONG_COLUMNS order. Returns the next free row.

    The measured error goes in its own column (sem or sd, per the panel's
    error_type); the other is a formula on sample_size, so it fills in as
    soon as n is typed.
    """
    r = start_row
    col = {name: first_col + i for i, name in enumerate(LONG_COLUMNS)}
    is_sd = result.panel.error_type == "sd"
    for bar in (items if items is not None else result.panel.bars):
        mean = bar.mean
        top = bar.mean_plus_err if bar.mean_plus_err is not None else mean
        err = None if mean is None or top is None else top - mean
        n_ref = f"{get_column_letter(col['sample_size'])}{r}"
        sem_ref = f"{get_column_letter(col['sem'])}{r}"
        sd_ref = f"{get_column_letter(col['sd'])}{r}"
        ws.cell(r, col["mean"], mean)
        if is_sd:
            ws.cell(r, col["sd"], err)
            ws.cell(r, col["sem"], f'=IF({n_ref}="","",{sd_ref}/SQRT({n_ref}))')
        else:
            ws.cell(r, col["sem"], err)
            ws.cell(r, col["sd"], f'=IF({n_ref}="","",{sem_ref}*SQRT({n_ref}))')
        ws.cell(r, col["treatment_group"], bar.label)
        if bar.n:
            ws.cell(r, col["sample_size"], bar.n)
        if result.paper_id is not None:
            ws.cell(r, col["paper_id"], result.paper_id)
        ws.cell(r, col["mean+sem"], top)
        if bar.uncertain:
            for c in range(first_col, first_col + len(LONG_COLUMNS)):
                ws.cell(r, c).fill = FLAG_FILL
        r += 1
    return r


def build_workbook(results: List[PanelResult]) -> Workbook:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["source", "figure", "panel", "y-axis title", "chart", "error bars", "groups",
                    "ticks used", "flagged", "warnings", "sheet"])
    for c in summary[1]:
        c.font = HEADER_FONT

    # One combined sheet: context columns, then the 12 long-format columns,
    # then the box-plot statistics (box rows only). Panels may mix SEM, SD
    # and box estimates, so the last long column is "mean+error" here.
    combined = wb.create_sheet("All data")
    combined.append(["source", "figure", "panel", "y-axis title", "error bars", "check"]
                    + LONG_COLUMNS[:-1] + ["mean+error"] + BOX_COLUMNS + ["x"])
    for c in combined[1]:
        c.font = HEADER_FONT
    combined_stats = 7 + len(LONG_COLUMNS)

    used = {"Summary", "All data"}
    next_combined = 2
    for res in results:
        p = res.panel
        is_box = p.kind in ("box", "violin")
        title = _sheet_title(f"{res.name} {p.y_title}".strip(), used)
        ws = wb.create_sheet(title)
        check_col = len(LONG_COLUMNS) + 2
        rows = line_rows(p) if p.kind == "line" else None
        if is_box:
            ws.append(long_header("sd") + ["", "check"] + BOX_COLUMNS)
            _write_box_rows(ws, 2, res, 1, check_col + 1)
            checks = [_box_check(bx) for bx in p.boxes]
        elif rows is not None:
            ws.append(long_header(p.error_type) + ["", "check", "series", "x"])
            _write_long_rows(ws, 2, res, items=rows)
            for i, row in enumerate(rows):
                ws.cell(2 + i, check_col + 1, row.label)
                ws.cell(2 + i, check_col + 2, row.x)
            checks = ["check point" if row.uncertain else "" for row in rows]
        else:
            ws.append(long_header(p.error_type) + ["", "check"])
            _write_long_rows(ws, 2, res)
            checks = ["check error bar" if bar.uncertain else "" for bar in p.bars]
        for c in ws[1]:
            c.font = HEADER_FONT
        for i, note in enumerate(checks):
            if note:
                ws.cell(2 + i, check_col, note)
        ws.column_dimensions["F"].width = 28

        error_label = f"{p.kind} (estimated)" if is_box else p.error_type.upper()
        for i, note in enumerate(checks):
            row = next_combined + i
            combined.cell(row, 1, res.source)
            combined.cell(row, 2, res.figure)
            combined.cell(row, 3, p.letter)
            combined.cell(row, 4, p.y_title)
            combined.cell(row, 5, error_label)
            combined.cell(row, 6, note)
        if is_box:
            next_combined = _write_box_rows(combined, next_combined, res, 7, combined_stats)
        elif rows is not None:
            for i, row in enumerate(rows):
                combined.cell(next_combined + i, combined_stats + len(BOX_COLUMNS), row.x)
            next_combined = _write_long_rows(combined, next_combined, res, first_col=7, items=rows)
        else:
            next_combined = _write_long_rows(combined, next_combined, res, first_col=7)

        summary.append([res.source, res.figure, p.letter, p.y_title, p.kind, error_label,
                        len(p.items), f"{p.calibration_inliers}/{len(p.ticks)}",
                        sum(1 for note in checks if note), "; ".join(p.warnings), title])

    for ws in (summary, combined):
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16
    return wb


def draw_overlay(image: np.ndarray, panels: List[Panel]) -> np.ndarray:
    """Annotated copy of the figure: axes, ticks, bars, caps, values. For eyeballing results."""
    out = image.copy()
    for p in panels:
        magenta, orange, blue = (200, 0, 200), (0, 150, 255), (220, 60, 0)
        cv2.line(out, (int(p.axis_x), p.plot_top), (int(p.axis_x), int(p.axis_y)), magenta, 2)
        cv2.line(out, (int(p.axis_x), int(p.axis_y)), (p.plot_right, int(p.axis_y)), magenta, 2)
        for ty, val, _ in p.ticks:
            color = (0, 170, 0) if val is not None else (0, 0, 255)
            cv2.circle(out, (int(p.axis_x) + 8, int(ty)), 4, color, -1)
        for b in p.boxes:
            colour = (0, 0, 255) if b.uncertain else orange
            cx = (b.x1 + b.x2) // 2
            cv2.rectangle(out, (b.x1, int(round(b.q3_px))), (b.x2, int(round(b.q1_px))), colour, 2)
            cv2.line(out, (b.x1, int(round(b.median_px))), (b.x2, int(round(b.median_px))), blue, 3)
            for y_end, y_box in ((b.hi_px, b.q3_px), (b.lo_px, b.q1_px)):
                cv2.line(out, (cx, int(round(y_box))), (cx, int(round(y_end))), colour, 1)
                cv2.line(out, (cx - 8, int(round(y_end))), (cx + 8, int(round(y_end))), colour, 2)
            if b.mean_px is not None:
                cv2.drawMarker(out, (cx, int(round(b.mean_px))), (0, 160, 0), cv2.MARKER_CROSS, 12, 2)
            for y in b.outliers_px:
                cv2.circle(out, (cx, int(round(y))), 4, magenta, 1)
        for s in p.series:
            # trace in a colour that stands out from the series' own
            pts = np.array([[int(x), int(y)] for x, y in s.trace], np.int32)
            if len(pts) > 1:
                cv2.polylines(out, [pts], False, magenta, 2)
            for pt in s.points:
                c = (int(round(pt.x_px)), int(round(pt.y_px)))
                cv2.circle(out, c, 5, orange, 2)
                for cap in (pt.cap_hi_px, pt.cap_lo_px):
                    if cap is not None:
                        cv2.line(out, (c[0] - 7, int(round(cap))), (c[0] + 7, int(round(cap))), blue, 2)
        for i, b in enumerate(p.bars, 1):
            if p.kind == "dot":
                cv2.line(out, (b.x1, int(round(b.top_px))), (b.x2, int(round(b.top_px))), orange, 3)
            else:
                cv2.rectangle(out, (b.x1, int(round(b.top_px))), (b.x2, int(p.axis_y)), orange, 2)
            cx = (b.x1 + b.x2) // 2
            cv2.putText(out, str(i), (cx - 5, int(p.axis_y) - 8), cv2.FONT_HERSHEY_SIMPLEX, 0.5, orange, 2)
            if b.cap_px is not None:
                cap_color = (0, 0, 255) if b.uncertain else blue
                cv2.line(out, (cx - 14, int(round(b.cap_px))), (cx + 14, int(round(b.cap_px))), cap_color, 3)
            if b.mean is not None:
                cv2.putText(out, f"{b.mean:.3g}", (b.x1, int(b.top_px) - 6),
                            cv2.FONT_HERSHEY_SIMPLEX, 0.55, blue, 2)
        if p.letter:
            cv2.putText(out, p.letter, (int(p.axis_x) + 10, p.plot_top + 20),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.8, magenta, 2)
    return out
